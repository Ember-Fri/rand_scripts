#!/usr/bin/env python3
"""
rfp_extract.py - Pull the RFP/IFB number, short description, and full Scope of
Work section out of each solicitation folder and write them to Excel.

Expected layout:
    <root>/
        RFP-123-26/
            RFP-123-26 (REQ 555).docx  <- name matched to the folder: parentheses
                                          ignored, 6T/6-Ton treated alike, fuzzy
                                          match as a fallback
        IFB-045-26/
            IFB-045-26.docx

Usage:
    python rfp_extract.py "C:\\path\\to\\Solicitations"
    python rfp_extract.py "C:\\path\\to\\Solicitations" -o rfp_summary.xlsx
    python rfp_extract.py "C:\\path\\to\\Solicitations" --no-prompt   # unattended

When no matching .docx is found, the script lists the folder's files and asks
you to pick one (or skip). Rows are flagged in the Notes column either way.

Requires:  pip install "python-docx>=1.1" openpyxl
"""

import argparse
import re
import sys
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

# ---------------------------------------------------------------------------
# Tweak these if your documents word things differently
# ---------------------------------------------------------------------------

# Section titles that count as the Scope of Work (regex, case-insensitive)
SCOPE_TITLES = [r"scope\s+of\s+work"]
# e.g. add r"scope\s+of\s+services", r"statement\s+of\s+work"

# Short-description label: "RFP Description:" (also accepts IFB/RFQ)
DESC_RE = re.compile(r"\b(?:RFP|IFB|RFQ)\s+Description\s*:\s*(.*)", re.I)

# RFP/IFB number in the folder title, e.g. "RFP-101-26", "IFB 45-26 6T Trucks",
# "RFP No. 123-26". If no match, the whole folder title (minus parentheses) is used.
FOLDER_ID_RE = re.compile(
    r"\b(?:RFP|IFB|RFQ)[\s\-_#:.]*(?:(?:No\.?|Number|#)[\s\-_#:.]*)?"
    r"[A-Za-z0-9]*\d[A-Za-z0-9\-./]*",
    re.I,
)

# Spellings treated as identical when matching file names to folder names
# (regex, replacement), applied after lowercasing. Add your own as you find them.
NAME_EQUIVALENTS = [
    (r"(\d+(?:\.\d+)?)\s*-?\s*(?:tons?|t)\b", r"\1t"),   # 6T, 6 T, 6-Ton, 6 tons -> 6t
    (r"\s*&\s*", " and "),
]

# How similar (0-1) a file name must be to the folder name to count as a match
FUZZY_THRESHOLD = 0.85

EXCEL_CELL_LIMIT = 32767

# ---------------------------------------------------------------------------
# Word auto-numbering (python-docx's .text drops list/heading numbers)
# ---------------------------------------------------------------------------


def _val(el, tag, default=None):
    child = el.find(qn(tag))
    return child.get(qn("w:val")) if child is not None else default


def _roman(n):
    out = ""
    for v, s in [(1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"),
                 (90, "xc"), (50, "l"), (40, "xl"), (10, "x"), (9, "ix"),
                 (5, "v"), (4, "iv"), (1, "i")]:
        while n >= v:
            out, n = out + s, n - v
    return out


def _letter(n):
    return chr(ord("a") + (n - 1) % 26) * ((n - 1) // 26 + 1)


def _fmt(n, fmt):
    return {
        "lowerLetter": _letter(n), "upperLetter": _letter(n).upper(),
        "lowerRoman": _roman(n), "upperRoman": _roman(n).upper(),
        "decimalZero": f"{n:02d}",
    }.get(fmt, str(n))


class Numberer:
    """Best-effort reconstruction of Word's automatic numbering."""

    def __init__(self, doc):
        self.nums = {}       # numId -> (abstractId, {ilvl: (fmt, lvlText, start)})
        self.counters = {}   # abstractId -> {ilvl: current value}
        try:
            numbering = doc.part.numbering_part.element
        except Exception:    # document has no numbering definitions
            return
        abstract = {}
        for an in numbering.findall(qn("w:abstractNum")):
            levels = {}
            for lvl in an.findall(qn("w:lvl")):
                levels[int(lvl.get(qn("w:ilvl")))] = (
                    _val(lvl, "w:numFmt", "decimal"),
                    _val(lvl, "w:lvlText", ""),
                    int(_val(lvl, "w:start", "1")),
                )
            abstract[an.get(qn("w:abstractNumId"))] = levels
        for num in numbering.findall(qn("w:num")):
            aid = _val(num, "w:abstractNumId")
            self.nums[num.get(qn("w:numId"))] = (aid, abstract.get(aid, {}))

    def _num_props(self, para):
        pPr = para._p.pPr
        numPr = pPr.find(qn("w:numPr")) if pPr is not None else None
        style = para.style
        while numPr is None and style is not None:      # numbering via style
            spPr = style.element.find(qn("w:pPr"))
            numPr = spPr.find(qn("w:numPr")) if spPr is not None else None
            style = style.base_style
        if numPr is None:
            return None
        num_id = _val(numPr, "w:numId")
        if num_id in (None, "0") or num_id not in self.nums:
            return None
        return num_id, int(_val(numPr, "w:ilvl", "0"))

    def label(self, para):
        """Return (rendered label like '3.2', (abstractId, ilvl)) or ('', None)."""
        props = self._num_props(para)
        if not props:
            return "", None
        num_id, ilvl = props
        aid, levels = self.nums[num_id]
        if ilvl not in levels:
            return "", (aid, ilvl)
        c = self.counters.setdefault(aid, {})
        c[ilvl] = c.get(ilvl, levels[ilvl][2] - 1) + 1
        for deeper in [k for k in c if k > ilvl]:
            del c[deeper]
        fmt, text, _ = levels[ilvl]
        if fmt == "bullet":
            return "\u2022", (aid, ilvl)
        if fmt == "none":
            return "", (aid, ilvl)

        def sub(m):
            k = int(m.group(1)) - 1
            lv = levels.get(k, ("decimal", "", 1))
            return _fmt(c.get(k, lv[2]), lv[0])

        return re.sub(r"%(\d)", sub, text), (aid, ilvl)


# ---------------------------------------------------------------------------
# Flatten the document into ordered lines
# ---------------------------------------------------------------------------


@dataclass
class Line:
    text: str
    is_table: bool = False
    num: tuple = None        # (abstractId, ilvl) if auto-numbered
    heading: int = None      # heading/outline level if any
    is_toc: bool = False


def _heading_level(para):
    name = (para.style.name or "").lower() if para.style is not None else ""
    m = re.match(r"heading\s*(\d)", name)
    if m:
        return int(m.group(1))
    pPr = para._p.pPr
    if pPr is not None:
        lvl = _val(pPr, "w:outlineLvl")
        if lvl is not None and int(lvl) < 9:
            return int(lvl) + 1
    return None


def _is_toc(para, text):
    name = (para.style.name or "").lower() if para.style is not None else ""
    return name.startswith("toc") or bool(
        re.search(r"(\t|\.{3,}|\s{3,})\s*\d+\s*$", text))


def _table_rows(table):
    """Yield one string per row, cells separated by ' | ' (merged cells deduped)."""
    for row in table.rows:
        seen, cells = [], []
        for cell in row.cells:
            if any(cell._tc is tc for tc in seen):   # merged cell repeats
                continue
            seen.append(cell._tc)
            t = " ".join(cell.text.split())
            if t:
                cells.append(t)
        if cells:
            yield " | ".join(cells)


def _table_cell_texts(table):
    seen = []
    for row in table.rows:
        for cell in row.cells:
            if any(cell._tc is tc for tc in seen):   # merged cell repeats
                continue
            seen.append(cell._tc)
            for p in cell.paragraphs:
                yield p.text.strip()
            for inner in cell.tables:
                yield from _table_cell_texts(inner)


def flatten(doc):
    """Return (section_lines, search_texts).

    section_lines: body-order lines used to cut out the Scope of Work.
    search_texts:  every paragraph incl. table cells and headers, used to find
                   the RFP number and description.
    """
    numberer = Numberer(doc)
    lines, search = [], []

    for sec in doc.sections:                       # headers often hold the RFP #
        for part in (sec.header, sec.first_page_header):
            try:
                search.extend(p.text.strip() for p in part.paragraphs)
            except Exception:
                pass

    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph):
            label, num = numberer.label(block)
            text = block.text.strip()
            full = f"{label} {text}" if label and text else text
            lines.append(Line(full, num=num, heading=_heading_level(block),
                              is_toc=_is_toc(block, text)))
            search.append(text)
        elif isinstance(block, Table):
            for row_text in _table_rows(block):
                lines.append(Line(row_text, is_table=True))
            search.extend(_table_cell_texts(block))
    return lines, [s for s in search if s]


# ---------------------------------------------------------------------------
# Field extraction
# ---------------------------------------------------------------------------

SCOPE_RE = re.compile(
    r"^\s*(?:(?:section|article|part)\s+)?(?P<num>\d{1,2}(?:\.\d+)*)?\.?\s*"
    r"[)\-\u2013\u2014:.]?\s*(?:" + "|".join(SCOPE_TITLES) + r")\b",
    re.I,
)
LEAD_NUM_RE = re.compile(
    r"^\s*(?:(?:section|article|part)\s+)?(\d{1,2}(?:\.\d+)*)(?:\.|\)|\t|\s)\s*\S", re.I)


def _num_tuple(s):
    parts = [int(p) for p in s.split(".") if p != ""]
    while len(parts) > 1 and parts[-1] == 0:        # "3.0" == "3"
        parts.pop()
    return tuple(parts)


def extract_scope(lines):
    start = None
    for i, ln in enumerate(lines):
        if ln.is_table or ln.is_toc or not ln.text:
            continue
        m = SCOPE_RE.match(ln.text)
        if m and (m.group("num") or ln.num or ln.heading):
            start, head_num = i, m.group("num")
            break
    if start is None:
        return "", "Scope of Work section not found"

    head = lines[start]
    head_tuple = _num_tuple(head_num) if head_num else None
    end = len(lines)
    for j in range(start + 1, len(lines)):
        ln = lines[j]
        if ln.is_table or ln.is_toc or not ln.text:
            continue
        # Same auto-numbered list, same or higher level -> next section
        if head.num and ln.num and ln.num[0] == head.num[0] and ln.num[1] <= head.num[1]:
            end = j
            break
        # Heading style at same or higher level -> next section
        if head.heading and ln.heading and ln.heading <= head.heading:
            end = j
            break
        # Typed number that comes after the Scope's number, e.g. "4." after "3."
        if head_tuple and len(ln.text) <= 120:
            m = LEAD_NUM_RE.match(ln.text)
            if m:
                t = _num_tuple(m.group(1))
                if len(t) <= len(head_tuple) and t > head_tuple:
                    end = j
                    break

    text = "\n".join(ln.text for ln in lines[start:end] if ln.text)
    note = "" if end < len(lines) else "Scope of Work ran to end of document; check for trailing content"
    return text, note


def extract_description(search):
    for i, t in enumerate(search):
        m = DESC_RE.search(t)
        if m:
            value = m.group(1).strip()
            if value:
                return value
            for nxt in search[i + 1:i + 4]:          # value on the next line/cell
                if nxt:
                    return nxt
    return ""


# ---------------------------------------------------------------------------
# Folder walking and Excel output
# ---------------------------------------------------------------------------


def strip_parens(name):
    """'RFP-101-26 (REQ 12345) (final)' -> 'RFP-101-26'"""
    name = re.sub(r"\([^)]*\)", " ", name)
    name = name.replace("(", " ").replace(")", " ")
    return re.sub(r"\s+", " ", name).strip(" -_")


def normalize_name(name):
    """Comparable form: no parentheses, equivalent spellings unified, no punctuation."""
    name = strip_parens(name).lower()
    for pattern, repl in NAME_EQUIVALENTS:
        name = re.sub(pattern, repl, name)
    return re.sub(r"[^a-z0-9]", "", name)


def rfp_number_from_folder(folder_name):
    clean_name = strip_parens(folder_name)
    m = FOLDER_ID_RE.search(clean_name)
    return m.group(0).strip(" -_.") if m else clean_name


def find_docx(folder):
    """Return (path or None, note, candidates). candidates = [(score, path)], best first."""
    docs = sorted((p for p in folder.rglob("*.docx") if not p.name.startswith("~$")),
                  key=lambda p: len(p.parts))
    if not docs:
        return None, "No .docx files in folder", []
    want = normalize_name(folder.name)

    for p in docs:                                   # exact match after normalizing
        if normalize_name(p.stem) == want:
            return p, "", []

    scored = sorted(((SequenceMatcher(None, want, normalize_name(p.stem)).ratio(), p)
                     for p in docs), key=lambda x: x[0], reverse=True)
    best_score, best = scored[0]
    if best_score >= FUZZY_THRESHOLD:
        note = f"Fuzzy name match ({best_score:.0%}): {best.name}"
        if len(scored) > 1 and scored[1][0] >= best_score - 0.03:
            note += f"; close runner-up {scored[1][1].name} ({scored[1][0]:.0%}), verify"
        return best, note, scored
    if len(docs) == 1:
        return best, f"Weak name match ({best_score:.0%}); used only .docx in folder: {best.name}", scored
    return None, (f"No matching .docx; closest was {best.name} ({best_score:.0%}) "
                  f"out of {len(docs)} files"), scored


# Manual selection is on unless --no-prompt is passed, input isn't a terminal,
# or you answer "q" at a prompt.
PROMPT_ENABLED = True


def choose_file(folder, candidates):
    """List the folder's files and let the user pick a .docx. Returns a path or None."""
    global PROMPT_ENABLED
    selectable = {p for _, p in candidates}
    others = sorted(p for p in folder.rglob("*")
                    if p.is_file() and p not in selectable and not p.name.startswith("~$"))

    print(f"\n{'=' * 70}\nNo matching .docx found for folder: {folder.name}")
    if candidates:
        print("Word documents (closest name first):")
        for i, (score, p) in enumerate(candidates, 1):
            print(f"  {i:>3})  {p.relative_to(folder)}   [{score:.0%} name match]")
    if others:
        print("Other files (not selectable; save as .docx to use):")
        for p in others:
            print(f"         {p.relative_to(folder)}")
    if not candidates:
        print("  (no .docx files to choose from - skipping)")
        return None

    print("    s)  skip this folder        q)  skip all remaining prompts")
    while True:
        try:
            choice = input(f"Select a file [1-{len(candidates)}, s, q]: ").strip().lower()
        except EOFError:
            PROMPT_ENABLED = False
            return None
        if choice in ("s", ""):
            return None
        if choice == "q":
            PROMPT_ENABLED = False
            return None
        if choice.isdigit() and 1 <= int(choice) <= len(candidates):
            return candidates[int(choice) - 1][1]
        print("  Not a valid choice, try again.")


def process_folder(folder):
    row = {"RFP Number": rfp_number_from_folder(folder.name),
           "Short Description": "", "Scope of Work": "",
           "Folder": folder.name, "Notes": []}
    path, note, candidates = find_docx(folder)
    if path is None and PROMPT_ENABLED:
        path = choose_file(folder, candidates)
        if path is not None:
            note = f"Manually selected: {path.relative_to(folder)}"
        elif candidates:
            note += "; skipped at manual selection"
    if note:
        row["Notes"].append(note)
    if path is None:
        return row
    try:
        doc = Document(str(path))
    except Exception as e:
        row["Notes"].append(f"Could not open {path.name}: {e}")
        return row

    lines, search = flatten(doc)
    row["Short Description"] = extract_description(search)
    if not row["Short Description"]:
        row["Notes"].append("'RFP Description:' not found")
    row["Scope of Work"], scope_note = extract_scope(lines)
    if scope_note:
        row["Notes"].append(scope_note)
    return row


def clean(value, notes):
    value = ILLEGAL_CHARACTERS_RE.sub("", value or "")
    if len(value) > EXCEL_CELL_LIMIT:
        notes.append(f"Text truncated to Excel's {EXCEL_CELL_LIMIT:,}-character cell limit")
        value = value[:EXCEL_CELL_LIMIT - 20] + "\n[...truncated]"
    return value


def write_excel(rows, out_path):
    cols = ["RFP Number", "Short Description", "Scope of Work", "Folder", "Notes"]
    widths = [22, 45, 100, 28, 45]
    wb = Workbook()
    ws = wb.active
    ws.title = "RFPs"
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center")
    for r in rows:
        notes = r["Notes"]
        values = [clean(r[c], notes) for c in cols[:-1]]
        values.append(clean("; ".join(notes), []))
        ws.append(values)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(out_path)


def main():
    ap = argparse.ArgumentParser(description="Extract RFP/IFB info into Excel.")
    ap.add_argument("root", type=Path, help="Folder containing the RFP/IFB folders")
    ap.add_argument("-o", "--output", type=Path, default=Path("rfp_summary.xlsx"))
    ap.add_argument("--no-prompt", action="store_true",
                    help="Don't ask to pick a file when no match is found; just flag the row")
    args = ap.parse_args()

    global PROMPT_ENABLED
    PROMPT_ENABLED = not args.no_prompt and sys.stdin.isatty()

    if not args.root.is_dir():
        sys.exit(f"Not a directory: {args.root}")

    folders = sorted(p for p in args.root.iterdir() if p.is_dir())
    rows = []
    for folder in folders:
        try:
            row = process_folder(folder)
        except Exception as e:                      # never let one file stop the run
            row = {"RFP Number": rfp_number_from_folder(folder.name), "Short Description": "", "Scope of Work": "",
                   "Folder": folder.name, "Notes": [f"Error: {e}"]}
        status = "OK" if not row["Notes"] else "CHECK: " + "; ".join(row["Notes"])
        print(f"{folder.name}: {status}")
        rows.append(row)

    write_excel(rows, args.output)
    flagged = sum(1 for r in rows if r["Notes"])
    print(f"\nWrote {len(rows)} rows to {args.output} ({flagged} flagged in Notes column)")


if __name__ == "__main__":
    main()
